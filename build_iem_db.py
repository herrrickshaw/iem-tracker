#!/usr/bin/env python3
"""
Build iem.duckdb from the DPIIT Industrial Entrepreneurs Memorandum (IEM)
materials in ~/Downloads and Dropbox.

Tables:
  iem_documents        one row per monthly Part A / Part B list PDF (679), joined to
                       the catalogue metadata, the local .md text, and the text length
  iem_dropbox_copies   the 2024 "Detailed Report" PDFs found in Dropbox, with md5 and
                       a canonical flag (byte-identical copies collapsed)
  annual_states        state-wise IEM counts and investment, 2021-2025 (Annual Report 2025-26)
  annual_sectors       sector-wise IEM counts and investment, Part A and Part B
  annual_trends        all-India IEM trend rows, 2021-2025
  catalogue_aggregates aggregate-statistics documents listed in the DPIIT catalogue
  part_a_records       record-level rows parsed from the Part A lists where the IEM-number
                       anchor matches (format-dependent; coverage per document is in
                       iem_documents.anchor_records)
"""
import glob
import hashlib
import os
import re
from pathlib import Path

import duckdb
import openpyxl

BASE = Path(__file__).resolve().parent
DB = BASE / "iem.duckdb"
CATALOGUE = Path.home() / "Downloads" / "DPIIT_IEM_PartA_PartB_catalogue_2026-10-07.xlsx"
ANALYSIS = Path.home() / "Dropbox" / "IEM_State_Sector_Analysis.xlsx"
DROPBOX_IEM_DIR = Path.home() / "Dropbox" / "Zotero_Unsorted_Bulk_Files"

ANCHOR = re.compile(r"(\d{1,4})\s*/\s*N\s*/\s*SIA\s*/\s*IM\s*O\s*/\s*(\d{4})\s*,?\s*(\d{2}/\d{2}/\d{4})")


def md5_of(path):
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def load_catalogue_rows(wb, sheet, part):
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    out = []
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        rec = dict(zip(header, r))
        out.append({
            "part": part,
            "period_from": str(rec.get("Period from")),
            "period_to": str(rec.get("Period to")),
            "granularity": rec.get("Granularity"),
            "dpiit_title": rec.get("DPIIT title"),
            "dpiit_category": rec.get("DPIIT category"),
            "pages": rec.get("Pages"),
            "size_kb": rec.get("Size (KB)"),
            "pdf_link": rec.get("PDF link"),
            "duplicate_of_dpiit_id": rec.get("Duplicate of DPIIT id"),
            "dpiit_doc_id": rec.get("DPIIT doc id"),
            "notes": rec.get("Notes"),
        })
    return out


def parse_wide_annual(ws, label_col_name):
    """States/Sectors sheets: a label column then (No., Inv) pairs per year, then a total."""
    rows = list(ws.iter_rows(values_only=True))
    years = [2021, 2022, 2023, 2024, 2025]
    data = []
    for r in rows[3:]:
        label = r[0]
        if not label or str(label).strip().lower() in ("total", "grand total", ""):
            continue
        for i, y in enumerate(years):
            n, inv = r[1 + 2 * i], r[2 + 2 * i]
            if n is None and inv is None:
                continue
            data.append((label, y, n, inv))
    return data


def main():
    if DB.exists():
        DB.unlink()
    con = duckdb.connect(str(DB))

    # ── 1. iem_documents ──────────────────────────────────────────────────
    wb = openpyxl.load_workbook(CATALOGUE, read_only=True, data_only=True)
    cat_rows = load_catalogue_rows(wb, "Part A – company lists", "A") + \
               load_catalogue_rows(wb, "Part B – company lists", "B")
    con.execute("""
        CREATE TABLE iem_documents (
            part VARCHAR, period_from VARCHAR, period_to VARCHAR, granularity VARCHAR,
            dpiit_title VARCHAR, dpiit_category VARCHAR, pages INTEGER, size_kb INTEGER,
            pdf_link VARCHAR, duplicate_of_dpiit_id VARCHAR, dpiit_doc_id BIGINT, notes VARCHAR,
            local_pdf VARCHAR, md_path VARCHAR, md_chars BIGINT, anchor_records INTEGER
        )
    """)
    local = {}
    for part in ("PartA", "PartB"):
        for p in glob.glob(str(Path.home() / "Downloads" / "DPIIT_IEM_lists" / part / "*.pdf")):
            m = re.search(r"_(\d+)\.pdf$", p)
            if m:
                local[m.group(1)] = (p, part)
    inserted = 0
    for r in cat_rows:
        docid = str(r["dpiit_doc_id"]) if r["dpiit_doc_id"] is not None else None
        pdf, md, chars, anchors = None, None, None, None
        if docid and docid in local:
            pdf, pt = local[docid]
            md = str(BASE / "md" / pt / (Path(pdf).name[:-4] + ".md"))
            if os.path.exists(md):
                text = open(md).read()
                chars = len(text)
                normalized = re.sub(r"\s+", " ", text)
                anchors = len(ANCHOR.findall(normalized))
        con.execute("""
            INSERT INTO iem_documents VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """, [r["part"], r["period_from"], r["period_to"], r["granularity"], r["dpiit_title"],
              r["dpiit_category"], r["pages"], r["size_kb"], r["pdf_link"],
              r["duplicate_of_dpiit_id"], int(docid) if docid else None, r["notes"],
              pdf, md, chars, anchors])
        inserted += 1
    print(f"iem_documents: {inserted} rows")

    # ── 2. iem_dropbox_copies (2024 detailed reports) ─────────────────────
    con.execute("""
        CREATE TABLE iem_dropbox_copies (
            file_name VARCHAR, md5 VARCHAR, size_bytes BIGINT, part VARCHAR,
            month VARCHAR, is_canonical BOOLEAN
        )
    """)
    seen = {}
    dropbox_rows = []
    for f in sorted(DROPBOX_IEM_DIR.glob("*IEM*Detailed*.pdf")):
        h = md5_of(f)
        part = "A" if "PartA" in f.name else "B"
        month = f.name.split("_IEM_")[0]
        dropbox_rows.append((f.name, h, f.stat().st_size, part, month))
    for name, h, size, part, month in dropbox_rows:
        canonical = h not in seen
        seen.setdefault(h, name)
        con.execute("INSERT INTO iem_dropbox_copies VALUES (?,?,?,?,?,?)",
                    [name, h, size, part, month, canonical])
    print(f"iem_dropbox_copies: {len(dropbox_rows)} files, "
          f"{sum(1 for r in dropbox_rows if r[1] not in seen or seen[r[1]] == r[0])} canonical")

    # ── 3. annual tables from the Dropbox analysis workbook ───────────────
    awb = openpyxl.load_workbook(ANALYSIS, read_only=True, data_only=True)
    con.execute("CREATE TABLE annual_states (part VARCHAR, state VARCHAR, year INTEGER, iem_count DOUBLE, investment_rs_cr DOUBLE)")
    con.execute("CREATE TABLE annual_sectors (part VARCHAR, sector VARCHAR, year INTEGER, iem_count DOUBLE, investment_rs_cr DOUBLE)")
    for sheet, part in (("States_PartA", "A"), ("States_PartB", "B")):
        for label, y, n, inv in parse_wide_annual(awb[sheet], "State"):
            con.execute("INSERT INTO annual_states VALUES (?,?,?,?,?)", [part, label, y, n, inv])
    for sheet, part in (("Sectors_PartA", "A"), ("Sectors_PartB", "B")):
        for label, y, n, inv in parse_wide_annual(awb[sheet], "Sector"):
            con.execute("INSERT INTO annual_sectors VALUES (?,?,?,?,?)", [part, label, y, n, inv])
    con.execute("CREATE TABLE annual_trends AS SELECT * FROM (VALUES (NULL)) t(x) WHERE false")
    con.execute("DROP TABLE annual_trends")
    con.execute("CREATE TABLE annual_trends (raw VARCHAR)")
    for r in awb["Trends"].iter_rows(values_only=True):
        if any(v is not None for v in r):
            con.execute("INSERT INTO annual_trends VALUES (?)", [" | ".join("" if v is None else str(v) for v in r)])
    print("annual_states:", con.execute("select count(*) from annual_states").fetchone()[0],
          "| annual_sectors:", con.execute("select count(*) from annual_sectors").fetchone()[0])

    # ── 4. catalogue aggregate statistics ─────────────────────────────────
    agg = list(wb["Aggregate statistics"].iter_rows(values_only=True))
    con.execute("CREATE TABLE catalogue_aggregates (part VARCHAR, header VARCHAR, dpiit_title VARCHAR, pages INTEGER, pdf_link VARCHAR, dpiit_doc_id BIGINT)")
    for r in agg[1:]:
        if not r or not r[0]:
            continue
        con.execute("INSERT INTO catalogue_aggregates VALUES (?,?,?,?,?,?)",
                    [r[0], r[1], r[2], r[4], r[6], r[8]])

    # ── 5. part_a_records (anchor-matched; format-dependent) ──────────────
    con.execute("CREATE TABLE part_a_records (iem_seq INTEGER, iem_year INTEGER, filing_date VARCHAR, source_doc VARCHAR)")
    for md in sorted(glob.glob(str(BASE / "md" / "PartA" / "*.md"))):
        normalized = re.sub(r"\s+", " ", open(md).read())
        for seq, year, date in ANCHOR.findall(normalized):
            con.execute("INSERT INTO part_a_records VALUES (?,?,?,?)",
                        [int(seq), int(year), date, Path(md).name])
    print("part_a_records:", con.execute("select count(*) from part_a_records").fetchone()[0])

    con.close()
    print(f"wrote {DB}")


if __name__ == "__main__":
    main()
